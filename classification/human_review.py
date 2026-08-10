"""Human review queue export, spreadsheet UX, and adjudication merge.

Builds a dedicated review artifact separating machine reference (coders A/B,
reviewer) from empty human adjudication columns. See classification/METHODOLOGY.md.
"""

from __future__ import annotations

import ast
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from classification.constants import (
    ALL_THEME_COLUMNS,
    CHARACTER_VALUES,
    HUMAN_CODER_A_COLUMNS,
    HUMAN_CODER_B_COLUMNS,
    HUMAN_CONTEXT_COLUMNS,
    HUMAN_EDITABLE_COLUMNS,
    HUMAN_QUEUE_COLUMNS,
    HUMAN_REVIEWER_COLUMNS,
    HUMAN_SUMMARY_COLUMNS,
    THEME_CODES,
    THEME_COLUMNS_A,
    THEME_COLUMNS_B,
    THEME_COLUMNS_HUMAN_MAIN,
    THEME_COLUMNS_HUMAN_QUEUE,
    THEME_COLUMNS_REV,
    THEME_DISPLAY_NAMES,
    TONE_VALUES,
)
from classification.schema import ThemeFlags


def _parse_emojis(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, list):
        return " ".join(str(item) for item in value)
    text = str(value).strip()
    if text.startswith("[") and text.endswith("]"):
        try:
            parsed = ast.literal_eval(text)
            if isinstance(parsed, list):
                return " ".join(str(item) for item in parsed)
        except (SyntaxError, ValueError):
            pass
    return text


def _flag_is_true(value: Any) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) == 1
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "x", "sim", "yes"}
    return bool(value)


def _themes_from_row(row: pd.Series, suffix: str) -> set[str]:
    affix = f"_{suffix}" if suffix else ""
    themes: set[str] = set()
    for code in THEME_CODES:
        col = f"theme_{code}{affix}"
        if col in row.index and _flag_is_true(row[col]):
            themes.add(code)
    na_col = f"theme_NA{affix}"
    if na_col in row.index and _flag_is_true(row[na_col]):
        themes.add("NA")
    return themes


def _themes_to_label(themes: set[str]) -> str:
    if not themes:
        return ""
    if themes == {"NA"}:
        return "N.A."
    ordered = [code for code in THEME_CODES if code in themes]
    return ",".join(ordered)


def _disagreement_summary(themes_a: set[str], themes_b: set[str], themes_rev: set[str]) -> str:
    parts: list[str] = []
    all_themes = themes_a | themes_b | themes_rev
    for code in sorted(all_themes - {"NA"}):
        holders = []
        if code in themes_a:
            holders.append("A")
        if code in themes_b:
            holders.append("B")
        if code in themes_rev:
            holders.append("rev")
        if len(holders) < 2 or set(holders) != {"A", "B", "rev"}:
            if len(holders) == 1:
                parts.append(f"{code}: só {holders[0]}")
            else:
                parts.append(f"{code}: {' vs '.join(holders)}")
    if "NA" in all_themes:
        na_holders = []
        if "NA" in themes_a:
            na_holders.append("A")
        if "NA" in themes_b:
            na_holders.append("B")
        if "NA" in themes_rev:
            na_holders.append("rev")
        if len(na_holders) < 3:
            parts.append(f"N.A.: {' vs '.join(na_holders)}")
    return "; ".join(parts)


def _bool_to_sheet(value: Any) -> int | str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return 1 if _flag_is_true(value) else 0


def initialize_human_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Ensure main-dataset human audit columns exist and are empty."""
    out = df.copy()
    for col in THEME_COLUMNS_HUMAN_MAIN:
        if col not in out.columns:
            out[col] = pd.NA
    for col in ("character_human", "tone_human", "human_notes", "reviewed_by", "reviewed_at"):
        if col not in out.columns:
            out[col] = pd.NA
    if "human_reviewed" not in out.columns:
        out["human_reviewed"] = False
    return out


def build_human_review_queue(df: pd.DataFrame) -> pd.DataFrame:
    """Build ordered human-review queue; no bare theme_* final columns."""
    base = initialize_human_columns(df)
    queue = base.loc[base["needs_human_review"].fillna(False).astype(bool)].copy()
    if queue.empty:
        return pd.DataFrame(columns=HUMAN_QUEUE_COLUMNS)

    rows: list[dict[str, Any]] = []
    for _, row in queue.iterrows():
        themes_a = _themes_from_row(row, "a")
        themes_b = _themes_from_row(row, "b")
        themes_rev = _themes_from_row(row, "")

        record: dict[str, Any] = {
            "hash_id": row.get("hash_id"),
            "case": row.get("case"),
            "comment_anon": row.get("comment_anon"),
            "emojis": _parse_emojis(row.get("emojis_list", row.get("emojis", ""))),
            "themes_a": _themes_to_label(themes_a),
            "themes_b": _themes_to_label(themes_b),
            "themes_rev": _themes_to_label(themes_rev),
            "disagreement": _disagreement_summary(themes_a, themes_b, themes_rev),
            "kappa_rev_a": row.get("kappa_reviewer_a"),
            "kappa_rev_b": row.get("kappa_reviewer_b"),
            "classification_source": row.get("classification_source"),
            "needs_human_review": row.get("needs_human_review"),
        }

        for col in THEME_COLUMNS_A:
            record[col] = _bool_to_sheet(row.get(col))
        record["character_a"] = row.get("character_a", "")
        record["tone_a"] = row.get("tone_a", "")
        record["confidence_a"] = row.get("confidence_a", "")

        for col in THEME_COLUMNS_B:
            record[col] = _bool_to_sheet(row.get(col))
        record["character_b"] = row.get("character_b", "")
        record["tone_b"] = row.get("tone_b", "")
        record["confidence_b"] = row.get("confidence_b", "")

        for final_col, rev_col in zip(ALL_THEME_COLUMNS, THEME_COLUMNS_REV, strict=True):
            record[rev_col] = _bool_to_sheet(row.get(final_col))
        record["character_rev"] = row.get("character", "")
        record["tone_rev"] = row.get("tone", "")
        record["confidence_rev"] = row.get("theme_conf", "")

        for col in THEME_COLUMNS_HUMAN_QUEUE:
            record[col] = ""
        record["character_h"] = ""
        record["tone_h"] = ""
        record["human_notes"] = ""
        record["reviewed_by"] = ""
        record["reviewed_at"] = ""

        rows.append(record)

    out = pd.DataFrame(rows, columns=HUMAN_QUEUE_COLUMNS)
    kappa_max = out[["kappa_rev_a", "kappa_rev_b"]].max(axis=1, skipna=True)
    return (
        out.assign(_kappa_max=kappa_max)
        .sort_values(
            by=["_kappa_max", "case", "hash_id"],
            ascending=[True, True, True],
            kind="stable",
        )
        .drop(columns=["_kappa_max"])
    )


def _queue_row_has_human_adjudication(row: pd.Series) -> bool:
    theme_filled = any(str(row.get(col, "")).strip() != "" for col in THEME_COLUMNS_HUMAN_QUEUE)
    char_filled = str(row.get("character_h", "")).strip() != ""
    tone_filled = str(row.get("tone_h", "")).strip() != ""
    return theme_filled or char_filled or tone_filled


def _validate_human_theme_flags(row: pd.Series) -> ThemeFlags:
    flags = {
        col.removeprefix("theme_").removesuffix("_h"): _flag_is_true(row.get(col))
        for col in THEME_COLUMNS_HUMAN_QUEUE
    }
    theme_kwargs = {f"theme_{code}": flags.get(code, False) for code in THEME_CODES}
    theme_kwargs["theme_NA"] = flags.get("NA", False)
    return ThemeFlags(**theme_kwargs)


def apply_human_adjudication(
    classified_df: pd.DataFrame,
    queue_df: pd.DataFrame,
    *,
    strict: bool = True,
) -> pd.DataFrame:
    """Merge filled human queue rows back into the classified dataset."""
    out = initialize_human_columns(classified_df)
    if queue_df.empty:
        return out

    queue = queue_df.copy()
    if "hash_id" not in queue.columns:
        msg = "queue_df must include hash_id"
        raise ValueError(msg)

    eligible = out.loc[out["needs_human_review"].fillna(False).astype(bool), "hash_id"]
    eligible_ids = set(eligible.astype(str))

    for _, row in queue.iterrows():
        hash_id = str(row["hash_id"])
        if hash_id not in eligible_ids:
            if strict:
                msg = f"hash_id {hash_id} is not in the human review queue"
                raise ValueError(msg)
            continue
        if not _queue_row_has_human_adjudication(row):
            continue

        flags = _validate_human_theme_flags(row)
        idx = out.index[out["hash_id"].astype(str) == hash_id]
        if idx.empty:
            if strict:
                msg = f"hash_id {hash_id} not found in classified_df"
                raise ValueError(msg)
            continue

        for code in THEME_CODES:
            out.loc[idx, f"theme_{code}"] = getattr(flags, f"theme_{code}")
            out.loc[idx, f"theme_{code}_human"] = getattr(flags, f"theme_{code}")
        out.loc[idx, "theme_NA"] = flags.theme_NA
        out.loc[idx, "theme_NA_human"] = flags.theme_NA

        character_h = str(row.get("character_h", "")).strip()
        tone_h = str(row.get("tone_h", "")).strip()
        if character_h:
            if character_h not in CHARACTER_VALUES:
                msg = f"Invalid character_h '{character_h}' for {hash_id}"
                raise ValueError(msg)
            out.loc[idx, "character"] = character_h
            out.loc[idx, "character_human"] = character_h
        if tone_h:
            if tone_h not in TONE_VALUES:
                msg = f"Invalid tone_h '{tone_h}' for {hash_id}"
                raise ValueError(msg)
            out.loc[idx, "tone"] = tone_h
            out.loc[idx, "tone_human"] = tone_h

        out.loc[idx, "human_notes"] = row.get("human_notes", "")
        out.loc[idx, "reviewed_by"] = row.get("reviewed_by", "")
        reviewed_at = row.get("reviewed_at", "")
        if not reviewed_at or (isinstance(reviewed_at, float) and pd.isna(reviewed_at)):
            reviewed_at = datetime.now(UTC).isoformat(timespec="seconds")
        out.loc[idx, "reviewed_at"] = reviewed_at
        out.loc[idx, "human_reviewed"] = True
        out.loc[idx, "classification_source"] = "human"

    return out


def _column_zone(column: str) -> str:
    if column in HUMAN_CONTEXT_COLUMNS or column in HUMAN_SUMMARY_COLUMNS:
        return "context"
    if column in HUMAN_CODER_A_COLUMNS or column in HUMAN_CODER_B_COLUMNS:
        return "reference"
    if column in HUMAN_REVIEWER_COLUMNS:
        return "reviewer"
    if column in HUMAN_EDITABLE_COLUMNS:
        return "human"
    return "meta"


def _write_codebook_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Codebook")
    ws.append(["Código", "Descrição"])
    for code, label in THEME_DISPLAY_NAMES.items():
        ws.append([code, label])
    ws.append([])
    ws.append(["character_h", "Afetivo | Cognitivo"])
    ws.append(["tone_h", "Positivo | Negativo | Neutro"])
    ws.append(["theme_*_h", "1 = presente, 0 = ausente; máximo 3 temas C; N.A. exclusivo"])
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)


def _write_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Instruções")
    lines = [
        "Fila de revisão humana — identidade digital brasileira",
        "",
        "1. Leia comment_anon e emojis com o contexto do case.",
        "2. Compare themes_a, themes_b e themes_rev (somente leitura).",
        "3. Preencha theme_C1_h … theme_NA_h com 0 ou 1.",
        "4. Preencha character_h e tone_h quando aplicável.",
        "5. Use human_notes para justificativas opcionais.",
        "6. Não edite colunas _a, _b ou _rev — são referência automática.",
        "7. Após revisar, salve o CSV e execute apply_human_adjudication no notebook.",
        "",
        "Critério de inclusão na fila: max(kappa_rev_a, kappa_rev_b) < 0,80.",
        "Ou seja, o revisor automático não alinhou bem com A nem com B.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 100


def export_human_review_queue(
    df: pd.DataFrame,
    path_csv: str | Path,
    path_xlsx: str | Path | None = None,
) -> pd.DataFrame:
    """Export human review queue as CSV and optional formatted Excel workbook."""
    queue = build_human_review_queue(df)
    csv_path = Path(path_csv)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    queue.to_csv(csv_path, index=False)

    if path_xlsx is not None:
        _export_human_review_xlsx(queue, Path(path_xlsx))

    return queue


def _export_human_review_xlsx(queue: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fila"

    fills = {
        "context": PatternFill("solid", fgColor="E8F4FD"),
        "reference": PatternFill("solid", fgColor="F2F2F2"),
        "reviewer": PatternFill("solid", fgColor="EDEDED"),
        "human": PatternFill("solid", fgColor="FFF4CC"),
        "meta": PatternFill("solid", fgColor="F9F9F9"),
    }

    headers = list(queue.columns) if not queue.empty else HUMAN_QUEUE_COLUMNS
    ws.append(headers)
    for cell in ws[1]:
        zone = _column_zone(str(cell.value))
        cell.fill = fills.get(zone, fills["meta"])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for _, row in queue.iterrows():
        ws.append([row.get(col, "") for col in headers])

    ws.freeze_panes = "E2"
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if header == "comment_anon":
            ws.column_dimensions[letter].width = 48
        elif header in {"disagreement", "human_notes"}:
            ws.column_dimensions[letter].width = 36
        elif header in HUMAN_CONTEXT_COLUMNS or header in HUMAN_SUMMARY_COLUMNS:
            ws.column_dimensions[letter].width = 18
        else:
            ws.column_dimensions[letter].width = 12

    if not queue.empty:
        last_row = len(queue) + 1
        theme_validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
        ws.add_data_validation(theme_validation)
        for col_name in THEME_COLUMNS_HUMAN_QUEUE:
            col_idx = headers.index(col_name) + 1
            letter = get_column_letter(col_idx)
            theme_validation.add(f"{letter}2:{letter}{last_row}")

        char_idx = headers.index("character_h") + 1
        char_letter = get_column_letter(char_idx)
        char_validation = DataValidation(
            type="list",
            formula1=f'"{",".join(CHARACTER_VALUES)}"',
            allow_blank=True,
        )
        ws.add_data_validation(char_validation)
        char_validation.add(f"{char_letter}2:{char_letter}{last_row}")

        tone_idx = headers.index("tone_h") + 1
        tone_letter = get_column_letter(tone_idx)
        tone_validation = DataValidation(
            type="list",
            formula1=f'"{",".join(TONE_VALUES)}"',
            allow_blank=True,
        )
        ws.add_data_validation(tone_validation)
        tone_validation.add(f"{tone_letter}2:{tone_letter}{last_row}")

    _write_codebook_sheet(wb)
    _write_instructions_sheet(wb)
    wb.save(path)
